#!/usr/bin/env python3
"""Offline UIOWA-081 drafting kit. Link checks are not assessment judgments."""
from __future__ import annotations
import argparse
import copy
import hashlib
import html
import io
import json
import math
import re
import sys
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

SCHEMA = "uiowa-executive-summary/v1"
GROUPS = {"ESS", "RIS", "IAM"}
DIMENSIONS = {"software_development", "security", "deployment", "ai_readiness"}
KINDS = {"strength", "gap", "disputed", "unknown"}
PHASES = ("0-90 days", "90-180 days", "180+ days")
BANNER = "SYNTHETIC EXAMPLE / NOT A UNIVERSITY ASSESSMENT"

class InputError(ValueError):
    """Actionable input error; no partially validated summary is emitted."""

def require(ok, message):
    if not ok:
        raise InputError(message)

def digest(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

def strict_json(text):
    def pairs(items):
        result = {}
        for key, value in items:
            require(key not in result, f"duplicate JSON key: {key}")
            result[key] = value
        return result
    def bad_number(value):
        raise InputError(f"non-finite JSON number: {value}")
    try:
        return json.loads(text, object_pairs_hook=pairs, parse_constant=bad_number)
    except json.JSONDecodeError as exc:
        raise InputError(f"invalid JSON at line {exc.lineno}: {exc.msg}") from exc

def text(value, label):
    require(isinstance(value, str) and bool(value.strip()), f"{label}: nonempty text required")
    require(all(ord(c) >= 32 or c in "\n\t" for c in value), f"{label}: control character")
    return value

def index(rows, label):
    require(isinstance(rows, list), f"{label}: array required")
    result = {}
    for row in rows:
        require(isinstance(row, dict), f"{label}: each record must be an object")
        key = text(row.get("id"), f"{label}.id")
        require(re.fullmatch(r"[A-Z][A-Z0-9-]*", key) is not None, f"invalid ID: {key}")
        require(key not in result, f"duplicate {label} ID: {key}")
        result[key] = row
    return result

def refs(value, registry, label, empty=False):
    require(isinstance(value, list) and (empty or len(value) > 0), f"{label}: reference array required")
    require(all(isinstance(v, str) for v in value), f"{label}: string references required")
    require(len(value) == len(set(value)), f"{label}: repeated reference")
    for key in value:
        require(key in registry, f"{label}: unknown reference {key}")

def hours(value, label):
    if value is None:
        return
    require(isinstance(value, list) and len(value) == 2, f"{label}: [low, high] or null required")
    require(all(type(v) in (int, float) and math.isfinite(v) and v >= 0 for v in value), f"{label}: finite nonnegative hours required")
    require(value[0] <= value[1], f"{label}: reversed range")

def validate(data):
    require(isinstance(data, dict), "summary packet must be an object")
    require(data.get("schema") == SCHEMA, "unsupported summary schema")
    require(data.get("status") == "SYNTHETIC_EXAMPLE", "this published kit accepts synthetic examples only; use the editable template for engagement drafting")
    for key in ("title", "scope", "resource_note", "limitation"):
        text(data.get(key), key)
    try:
        date.fromisoformat(data["as_of"])
    except (KeyError, ValueError, TypeError):
        raise InputError("as_of: ISO calendar date required") from None
    sources = index(data.get("sources"), "sources")
    findings = index(data.get("findings"), "findings")
    recs = index(data.get("recommendations"), "recommendations")
    require(sources and findings and recs, "sources, findings and recommendations must be nonempty")
    require(len(set(sources) | set(findings) | set(recs)) == len(sources)+len(findings)+len(recs), "IDs must be unique across record types")
    for sid, source in sources.items():
        for key in ("title", "version", "locator", "excerpt"):
            text(source.get(key), f"{sid}.{key}")
        require(source.get("evidence_class") == "synthetic", f"{sid}: synthetic evidence label required")
        require(source.get("sha256") == digest(source["excerpt"]), f"{sid}: excerpt digest mismatch")
    for fid, finding in findings.items():
        for key in ("text", "limitation"):
            text(finding.get(key), f"{fid}.{key}")
        require(isinstance(finding.get("kind"), str) and finding["kind"] in KINDS, f"{fid}: unsupported finding kind")
        refs(finding.get("groups"), GROUPS, f"{fid}.groups")
        require(isinstance(finding.get("dimension"), str) and finding["dimension"] in DIMENSIONS, f"{fid}: unsupported dimension")
        refs(finding.get("source_ids"), sources, f"{fid}.source_ids")
    for rid, rec in recs.items():
        for key in ("title", "decision", "owner_role", "expected_outcome", "verification", "assumptions"):
            text(rec.get(key), f"{rid}.{key}")
        require(rec.get("phase") in PHASES, f"{rid}: unsupported phase")
        refs(rec.get("finding_ids"), findings, f"{rid}.finding_ids")
        refs(rec.get("depends_on", []), recs, f"{rid}.depends_on", empty=True)
        require(rid not in rec.get("depends_on", []), f"{rid}: self dependency")
        hours(rec.get("one_time_hours"), f"{rid}.one_time_hours")
        hours(rec.get("monthly_hours"), f"{rid}.monthly_hours")
    visited, active = set(), set()
    def walk(rid):
        require(rid not in active, f"recommendation dependency cycle at {rid}")
        if rid in visited:
            return
        active.add(rid)
        for dep in recs[rid].get("depends_on", []):
            require(PHASES.index(recs[dep]["phase"]) <= PHASES.index(recs[rid]["phase"]), f"{rid}: dependency {dep} is in a later phase")
            walk(dep)
        active.remove(rid)
        visited.add(rid)
    for rid in recs:
        walk(rid)
    brief = data.get("decision_brief")
    require(isinstance(brief, dict), "decision_brief: object required")
    text(brief.get("text"), "decision_brief.text")
    refs(brief.get("finding_ids"), findings, "decision_brief.finding_ids")
    refs(brief.get("recommendation_ids"), recs, "decision_brief.recommendation_ids")
    require("maturity_score" not in data, "the summary does not compute an institutional maturity score")
    return sources, findings, recs

def inspect_handoff(data):
    """Preserve workbench context. Notes never become findings automatically."""
    require(isinstance(data, dict), "handoff must be an object")
    require(data.get("schema") == "uiowa-rfq18649-analyst-handoff-draft/v1", "unsupported handoff schema")
    require(data.get("status") == "DRAFT_NON_AUTHORITATIVE", "expected a draft workbench handoff")
    require(data.get("report_mode") == "UNTRUSTED_INSPECTION", "handoff: expected untrusted inspection mode")
    receipt = data.get("report_receipt_sha256")
    require(isinstance(receipt, str) and re.fullmatch("[0-9a-f]{64}", receipt) is not None, "handoff: invalid report receipt")
    rows = data.get("cell_notes")
    require(isinstance(rows, list) and len(rows) == 12, "handoff: expected 12 cells")
    cells = set()
    normalized = copy.deepcopy(data)
    for row in normalized["cell_notes"]:
        require(isinstance(row, dict), "handoff: invalid cell")
        group, dimension = row.get("group"), row.get("dimension")
        dimension = "software_development" if dimension == "software" else dimension
        require(isinstance(group, str) and isinstance(dimension, str) and group in GROUPS and dimension in DIMENSIONS, "handoff: unknown group or dimension")
        require((group, dimension) not in cells, "handoff: duplicate cell")
        cells.add((group, dimension))
        row["dimension"] = dimension
        text(row.get("compiler_status"), "handoff.compiler_status")
        text(row.get("disposition"), "handoff.disposition")
        require(isinstance(row.get("analyst_note"), str), "handoff: analyst_note must be text")
    require(isinstance(data.get("authority"), dict) and data["authority"].get("current_evidence_review_authority") is False, "handoff is not a non-authoritative draft")
    return {"status": "DRAFT_CONTEXT_ONLY", "source_sha256": digest(json.dumps(data, sort_keys=True, ensure_ascii=False)), "handoff": normalized,
            "interpretation": "Original compiler statuses, notes and authority fields are retained. No notes, approvals or numeric ratings are promoted into findings."}

def amount(value):
    return "UNKNOWN (estimate required)" if value is None else f"{value[0]:g}-{value[1]:g} person-hours"

def citation(ids):
    return "[" + ", ".join(ids) + "]"

def blocks(data):
    sources, findings, recs = validate(data)
    out = [("h1", data["title"]), ("notice", BANNER + " | " + data["as_of"]), ("p", data["scope"]), ("h2", "Decision brief")]
    brief = data["decision_brief"]
    out.append(("p", brief["text"] + " " + citation(brief["finding_ids"] + brief["recommendation_ids"])))
    out.append(("h2", "Strengths to retain"))
    for fid, f in findings.items():
        if f["kind"] == "strength":
            out.append(("p", f["text"] + " " + citation([fid])))
    out.append(("h2", "Consequential gaps and uncertainty"))
    for fid, f in findings.items():
        if f["kind"] != "strength":
            out.append(("p", f"{f['kind'].upper()}: {f['text']} {citation([fid])}"))
    out.extend([("p", "Interpretation limit: " + data["limitation"]), ("break", ""), ("h2", "Decisions and practical sequence")])
    for rid, rec in recs.items():
        out.extend([("h3", f"{rid} | {rec['phase']} | {rec['title']}"),
                    ("p", f"Decision: {rec['decision']} Owner role: {rec['owner_role']}. {citation(rec['finding_ids'])}"),
                    ("p", f"Planning effort: {amount(rec['one_time_hours'])} one time; {amount(rec['monthly_hours'])} per month. Dependencies: {', '.join(rec.get('depends_on', [])) or 'none in this example'}."),
                    ("p", f"Proposed outcome: {rec['expected_outcome']} Check: {rec['verification']}"),
                    ("small", "Assumptions: " + rec["assumptions"])])
    out.extend([("h2", "Resource and decision boundaries"), ("p", data["resource_note"]),
                ("small", "Support: traceability.md maps every finding and recommendation to versioned synthetic excerpts. Structural checks do not establish semantic adequacy, institutional maturity, approval, or achieved benefit.")])
    return out

def template_blocks():
    return [("h1", "Leadership executive summary"), ("notice", "EDITABLE TEMPLATE / NOT ASSESSED / NO INSTITUTIONAL FINDINGS"),
            ("p", "[Engagement and review version] | [Prepared date] | [Decision audience]"),
            ("h2", "Decision brief"), ("p", "[In 60-90 words: what leadership should retain, what to address next, why it matters, and which decision is requested. Cite finding and recommendation IDs. State the extent of evidence and the judgment's limits.]"),
            ("h2", "Scope and basis"), ("p", "[ESS / RIS / IAM; software development, security, deployment/operations and AI readiness. Add actual evidence dates, source types, sampled services and unassessed areas. Do not equate missing evidence with missing practice.]"),
            ("h2", "Strengths to retain"), ("p", "[Describe two concrete observed strengths, their service/group scope and supporting source-linked finding IDs. Explain the useful outcome; avoid endorsing every team or every release from a small sample.]"),
            ("h2", "Consequential gaps and uncertainty"), ("p", "[Describe two or three evidence-backed gaps or unresolved questions. Separate confirmed observations, conflicting sources and unknowns. Explain the operational consequence without inventing a probability or departmental score.]"),
            ("break", ""), ("h2", "Decisions and practical sequence"),
            ("p", "[R-ID | 0-90 days | accountable role | proposed action | linked F-IDs | prerequisite decision]"),
            ("p", "[R-ID | 90-180 days | accountable role | proposed action | linked F-IDs | dependency on earlier work]"),
            ("p", "[R-ID | 180+ days | accountable role | conditional next step | observable evidence needed before expansion]"),
            ("h2", "Resource implications"), ("p", "[Show one-time implementation/training person-hours separately from recurring monthly workload. State low/high assumptions, constrained roles, overlaps and unknown estimates. Capacity saved is not cash saved. No proposed allowance is an approved budget.]"),
            ("h2", "Expected outcomes and verification"), ("p", "[For each R-ID: proposed outcome, current baseline or UNKNOWN, observation window, data source, success measure, and interpretation caveat. Do not present the target as an achieved result or a causal effect.]"),
            ("h2", "Decisions still required"), ("p", "[Priorities and owner roles to confirm; unavailable evidence; unresolved interpretations; resource allocation. Preserve disagreement rather than silently choosing a preferred account.]"),
            ("small", "Review: trace each substantive statement to F/R/source IDs; retain source versions/locators; reconcile wording with the report; retain UNKNOWN, disputed and unassessed states. Drafting does not create maturity, submission, commercial or approval authority.")]

def markdown(items):
    return "\n\n".join(("# " if kind == "h1" else "## " if kind == "h2" else "### " if kind == "h3" else "> " if kind == "notice" else "") + value if kind != "break" else "---" for kind, value in items) + "\n"

def html_document(items):
    body = []
    for kind, value in items:
        if kind == "break":
            body.append('<div class="page-break"></div>')
        else:
            tag = kind if kind in ("h1", "h2", "h3") else "p"
            body.append(f'<{tag} class="{kind}">{html.escape(value)}</{tag}>')
    return '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Leadership executive summary</title><style>body{font:11pt/1.45 Arial,sans-serif;max-width:780px;margin:40px auto;padding:0 24px;color:#16242f}h1{font-size:24pt}h2{font-size:15pt;margin-top:1.2em}h3{font-size:11.5pt}.notice{font-weight:bold;border:2px solid;padding:10px}.small{font-size:9.5pt}p{overflow-wrap:anywhere}@media print{body{margin:0;max-width:none}.page-break{break-before:page}h2,h3{break-after:avoid}@page{size:letter;margin:0.65in}}</style></head><body><main>' + "\n".join(body) + '</main></body></html>\n'

def docx_bytes(items):
    """Optional python-docx export; stable metadata/ZIP timestamps, no macros."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise InputError("DOCX output requires python-docx; core Markdown/HTML/JSON remains dependency-free") from exc
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.header_distance = sec.footer_distance = Inches(0.25)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.color.rgb = RGBColor.from_string("16242F")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.04
    doc.styles["Title"].font.size = Pt(23)
    for name, size in (("Heading 1", 14), ("Heading 2", 11)):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    sec.header.paragraphs[0].text = "TJLabs | Assessment preparation | UIOWA-081"
    sec.header.paragraphs[0].style = doc.styles["Normal"]
    for run in sec.header.paragraphs[0].runs:
        run.font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.text = "SYNTHETIC / TEMPLATE - NOT ASSESSED                                      "
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    for kind, value in items:
        if kind == "break":
            doc.add_page_break(); continue
        p = doc.add_paragraph(value, {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2"}.get(kind, "Normal"))
        if kind in ("notice", "small"):
            for run in p.runs:
                run.font.size = Pt(8.5 if kind == "small" else 9)
                run.bold = kind == "notice"
        p.paragraph_format.widow_control = True
    doc.core_properties.author = "TJLabs / ZZ-INKSTONE-81Q"
    doc.core_properties.title = "UIOWA-081 leadership executive-summary preparation"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 9, 19, tzinfo=timezone.utc)
    raw = io.BytesIO(); doc.save(raw)
    # Prune unused styles while retaining all package parts and relationships.
    from lxml import etree
    original = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    result = io.BytesIO()
    used = {"Normal", "Title", "Heading1", "Heading2", "Heading3", "Header", "Footer", "DefaultParagraphFont", "TableNormal", "NoList"}
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(original.namelist()):
            payload = original.read(name)
            if name in ("word/styles.xml", "word/stylesWithEffects.xml"):
                root = etree.fromstring(payload)
                for node in list(root):
                    if node.tag == qn("w:latentStyles") or (node.tag == qn("w:style") and node.get(qn("w:styleId")) not in used):
                        root.remove(node)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            entry = zipfile.ZipInfo(name, (2026, 9, 19, 0, 0, 0)); entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, payload)
    return result.getvalue()

def trace(data):
    sources, findings, recs = validate(data)
    records = []
    for fid, f in findings.items():
        records.append({"statement_id": fid, "kind": f["kind"], "text": f["text"], "limitation": f["limitation"],
                        "source_ids": f["source_ids"], "sources": [copy.deepcopy(sources[sid]) for sid in f["source_ids"]]})
    for rid, r in recs.items():
        ids = sorted({sid for fid in r["finding_ids"] for sid in findings[fid]["source_ids"]})
        records.append({"statement_id": rid, "kind": "proposed_recommendation", "finding_ids": r["finding_ids"], "source_ids": ids,
                        "expected_outcome_status": "PROPOSED_NOT_MEASURED", "assumptions": r["assumptions"]})
    brief = data["decision_brief"]
    records.append({"statement_id": "DECISION-BRIEF", "kind": "editorial_synthesis", "finding_ids": brief["finding_ids"], "recommendation_ids": brief["recommendation_ids"], "text": brief["text"]})
    return {"status": "SYNTHETIC_DRAFT_TRACE", "input_sha256": digest(json.dumps(data, sort_keys=True, ensure_ascii=False)), "records": records,
            "limitation": "Structural referential integrity and excerpt consistency only. Human review must evaluate whether each source supports the wording and scope."}

def trace_markdown(data):
    sources, findings, recs = validate(data)
    lines = ["# Executive-summary support register", "> " + BANNER, "Structural links are not an automatic semantic review."]
    for fid, f in findings.items():
        lines += [f"## {fid} | {f['kind']}", f["text"], f"Scope: {', '.join(f['groups'])} / {f['dimension']}. Limitation: {f['limitation']}", "Support: " + citation(f["source_ids"])]
    for rid, r in recs.items():
        lines += [f"## {rid} | proposed recommendation", "Findings: " + citation(r["finding_ids"]), "Assumptions: " + r["assumptions"]]
    for sid, s in sources.items():
        lines += [f"## {sid} | {s['title']}", f"Version: {s['version']}. Locator: `{s['locator']}`. Excerpt SHA-256: `{s['sha256']}`.", "> " + s["excerpt"].replace("\n", "\n> ")]
    return "\n\n".join(lines) + "\n"

def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--docx", action="store_true")
    parser.add_argument("--handoff", type=Path)
    args = parser.parse_args(argv)
    try:
        data = strict_json(args.input.read_text(encoding="utf-8"))
        example, template = blocks(data), template_blocks()
        generated = {"executive-summary.md": markdown(example).encode(), "executive-summary.html": html_document(example).encode(),
                     "editable-template.md": markdown(template).encode(), "editable-template.html": html_document(template).encode(),
                     "traceability.json": (json.dumps(trace(data), indent=2, ensure_ascii=False) + "\n").encode(), "traceability.md": trace_markdown(data).encode()}
        if args.handoff:
            context = inspect_handoff(strict_json(args.handoff.read_text(encoding="utf-8")))
            generated["workbench-context.json"] = (json.dumps(context, indent=2, ensure_ascii=False) + "\n").encode()
        if args.docx:
            generated["executive-summary.docx"] = docx_bytes(example)
            generated["editable-template.docx"] = docx_bytes(template)
        require(not args.out.exists(), "output directory already exists; choose a new run directory to preserve earlier drafts")
        args.out.mkdir(parents=True)
        for name, content in generated.items():
            (args.out / name).write_bytes(content)
        manifest = {"status": "SYNTHETIC_PREPARATION_ONLY", "files": {k: hashlib.sha256(v).hexdigest() for k, v in sorted(generated.items())}}
        (args.out / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"result": "PASS_STRUCTURAL", "files": len(generated), "source_records": len(data['sources']), "findings": len(data['findings']), "recommendations": len(data['recommendations'])}))
        return 0
    except (InputError, OSError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

if __name__ == "__main__":
    raise SystemExit(main())
