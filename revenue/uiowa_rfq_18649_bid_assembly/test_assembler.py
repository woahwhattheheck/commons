"""Executable real-build and negative-control tests; all file mutations stay in temp dirs."""
import copy
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn
import assembler as a

ROOT = Path(__file__).resolve().parent


class AssemblyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.session = tempfile.TemporaryDirectory()
        cls.base = Path(cls.session.name)
        cls.pristine = cls.base / 'pristine'
        cls.receipt = a.build(ROOT / 'inputs/packet.json', cls.pristine)

    @classmethod
    def tearDownClass(cls):
        cls.session.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'inputs'
        shutil.copytree(ROOT / 'inputs', self.inputs)
        self.packet_path = self.inputs / 'packet.json'
        self.packet = json.loads(self.packet_path.read_text())
        self.folder = self.root / 'copy'
        shutil.copytree(self.pristine, self.folder)

    def write_packet(self):
        self.packet_path.write_text(json.dumps(self.packet))

    def rehash(self, filename):
        manifest = json.loads((self.folder / '07_MANIFEST.json').read_text())
        manifest['files'][filename] = a.digest((self.folder / filename).read_bytes())
        a.write_json(self.folder / '07_MANIFEST.json', manifest)

    def test_real_build_opens_all_ten_outputs(self):
        result = a.verify(self.folder)
        self.assertEqual(result['verified_files'], 10)
        self.assertEqual(result['pdf_pages'], 7)
        self.assertEqual(result['pdf_links'], 19)
        self.assertEqual(result['docx_bookmarks'], 6)
        self.assertEqual(result['docx_links'], 19)
        self.assertEqual(result['html_links'], 21)
        self.assertFalse(result['submission_authorized'])

    def test_actual_commercial_git_blob_matches(self):
        actual = (self.folder / '05_COMMERCIAL_SOURCE.md').read_bytes()
        self.assertEqual(len(actual), 2621)
        self.assertEqual(a.git_digest(actual), 'b6e9ca58984c15d96b497f3bb51000992fdb9b5f')
        self.assertEqual(actual, (ROOT / 'inputs/COMMERCIAL.md').read_bytes())

    def test_repeated_real_build_is_byte_identical(self):
        other = self.root / 'second'
        a.build(self.packet_path, other)
        self.assertEqual({f.name: a.digest(f.read_bytes()) for f in other.iterdir()},
                         {f.name: a.digest(f.read_bytes()) for f in self.folder.iterdir()})

    def test_original_source_text_survives_pdf(self):
        packet, data = a.load_packet(self.packet_path)
        with fitz.open(self.folder / '01_REVIEW_DRAFT.pdf') as pdf:
            rendered = ' '.join(pdf[2].get_text().split())
        for paragraph in data.decode().strip().split('\n\n'):
            self.assertIn(' '.join(a.plain(paragraph).split()), rendered)

    def test_source_hash_change_is_rejected(self):
        (self.inputs / 'COMMERCIAL.md').write_text('Fabricated source')
        with self.assertRaisesRegex(a.AssemblyError, 'SOURCE_HASH_MISMATCH'):
            a.load_packet(self.packet_path)

    def test_git_binding_independent_of_sha256(self):
        self.packet['source']['git_blob_sha'] = '0' * 40
        self.write_packet()
        with self.assertRaisesRegex(a.AssemblyError, 'SOURCE_HASH_MISMATCH'):
            a.load_packet(self.packet_path)

    def test_duplicate_json_key_rejected(self):
        with self.assertRaisesRegex(a.AssemblyError, 'duplicate'):
            a.strict_json('{"a": 1, "a": 2}')

    def test_nonfinite_json_rejected(self):
        for val in ('NaN', 'Infinity', '-Infinity'):
            with self.subTest(val=val), self.assertRaises(a.AssemblyError):
                a.strict_json('{"a": ' + val + '}')

    def test_boolean_schema_not_integer_schema(self):
        self.packet['schema'] = True
        self.write_packet()
        with self.assertRaises(a.AssemblyError): a.load_packet(self.packet_path)

    def test_unknown_packet_key_rejected(self):
        self.packet['approved'] = True
        self.write_packet()
        with self.assertRaises(a.AssemblyError): a.load_packet(self.packet_path)

    def test_traversal_and_nonportable_filenames_rejected(self):
        for value in ('../private', '/etc/passwd', 'a/b', 'a\\b', '..', 'x..pdf', 'x:y', '.hidden'):
            with self.subTest(value=value), self.assertRaises(a.AssemblyError): a.local_name(value)

    def test_symlink_source_rejected(self):
        name = self.inputs / 'COMMERCIAL.md'
        name.unlink(); name.symlink_to(ROOT / 'inputs/COMMERCIAL.md')
        with self.assertRaisesRegex(a.AssemblyError, 'symlink'): a.load_packet(self.packet_path)

    def test_duplicate_requirement_id_rejected(self):
        self.packet['requirements'].append(copy.deepcopy(self.packet['requirements'][0]))
        self.write_packet()
        with self.assertRaisesRegex(a.AssemblyError, 'duplicate'): a.load_packet(self.packet_path)

    def test_requirement_deletion_rejected(self):
        self.packet['requirements'].pop()
        self.write_packet()
        with self.assertRaisesRegex(a.AssemblyError, 'universe'): a.load_packet(self.packet_path)

    def test_requirement_approval_cannot_be_invented(self):
        self.packet['requirements'][0]['status'] = 'APPROVED'
        self.write_packet()
        with self.assertRaisesRegex(a.AssemblyError, 'authorizing'): a.load_packet(self.packet_path)

    def test_required_missing_fields_are_never_synthetic_evidence(self):
        index = json.loads((self.folder / '02_ATTACHMENT_INDEX.json').read_text())
        placeholders = [x for x in index['attachments'] if x['status'] == 'MISSING_PLACEHOLDER']
        self.assertEqual({x['filename'] for x in placeholders}, {'03_FINANCIALS_MISSING.txt', '04_QUALIFICATIONS_MISSING.txt'})
        self.assertIn('audited statements and annual reports', (self.folder / '03_FINANCIALS_MISSING.txt').read_text())
        self.assertFalse(index['submission_authorized'])

    def test_existing_output_directory_preserved(self):
        sentinel = self.folder / 'OWNER_DATA.txt'; sentinel.write_text('leave untouched')
        before = {p.name: p.read_bytes() for p in self.folder.iterdir()}
        with self.assertRaisesRegex(a.AssemblyError, 'new directory'): a.build(self.packet_path, self.folder)
        self.assertEqual(before, {p.name: p.read_bytes() for p in self.folder.iterdir()})

    def test_output_symlink_parent_rejected(self):
        link = self.root / 'link'; link.symlink_to(self.root, target_is_directory=True)
        with self.assertRaisesRegex(a.AssemblyError, 'symlink'): a.build(self.packet_path, link / 'output')

    def test_output_extra_file_fails_closed(self):
        (self.folder / 'foreign.txt').write_text('not expected')
        with self.assertRaisesRegex(a.AssemblyError, 'file-set'): a.verify(self.folder)

    def test_output_missing_file_fails_closed(self):
        (self.folder / '03_FINANCIALS_MISSING.txt').unlink()
        with self.assertRaisesRegex(a.AssemblyError, 'file-set'): a.verify(self.folder)

    def test_output_symlink_rejected(self):
        item = self.folder / '05_COMMERCIAL_SOURCE.md'; item.unlink()
        item.symlink_to(ROOT / 'inputs/COMMERCIAL.md')
        with self.assertRaisesRegex(a.AssemblyError, 'symlink'): a.verify(self.folder)

    def test_output_byte_tamper_rejected(self):
        with (self.folder / '03_FINANCIALS_MISSING.txt').open('a') as handle: handle.write('changed')
        with self.assertRaisesRegex(a.AssemblyError, 'OUTPUT_HASH_MISMATCH'): a.verify(self.folder)

    def test_rehashed_wrong_pdf_link_rejected(self):
        path = self.folder / '01_REVIEW_DRAFT.pdf'
        with fitz.open(path) as pdf:
            link = pdf[0].get_links()[0]; link['page'] = 6
            pdf[0].update_link(link); pdf.saveIncr()
        self.rehash(path.name)
        with self.assertRaisesRegex(a.AssemblyError, 'target mismatch'): a.verify(self.folder)

    def test_rehashed_wrong_docx_page_field_rejected(self):
        path = self.folder / '01_REVIEW_DRAFT.docx'
        doc = Document(path)
        for node in doc.element.iter(qn('w:fldSimple')):
            if node.get(qn('w:instr'), '').startswith('PAGEREF '):
                node.set(qn('w:instr'), 'PAGEREF qualifications \\h'); break
        doc.save(path); self.rehash(path.name)
        with self.assertRaisesRegex(a.AssemblyError, 'page-reference'): a.verify(self.folder)

    def test_rehashed_stale_docx_cached_page_rejected(self):
        path = self.folder / '01_REVIEW_DRAFT.docx'
        doc = Document(path)
        for node in doc.element.iter(qn('w:fldSimple')):
            if node.get(qn('w:instr'), '').startswith('PAGEREF '):
                next(node.iter(qn('w:t'))).text = '99'; break
        doc.save(path); self.rehash(path.name)
        with self.assertRaisesRegex(a.AssemblyError, 'cached page-reference'): a.verify(self.folder)

    def test_rehashed_wrong_docx_bookmark_rejected(self):
        path = self.folder / '01_REVIEW_DRAFT.docx'
        doc = Document(path)
        next(doc.element.iter(qn('w:bookmarkStart'))).set(qn('w:name'), 'unknown')
        doc.save(path); self.rehash(path.name)
        with self.assertRaisesRegex(a.AssemblyError, 'bookmark'): a.verify(self.folder)

    def test_rehashed_wrong_html_link_rejected(self):
        path = self.folder / '00_START_HERE.html'
        path.write_text(path.read_text().replace('href="#financials"', 'href="#wrong"'))
        self.rehash(path.name)
        with self.assertRaisesRegex(a.AssemblyError, 'HTML navigation'): a.verify(self.folder)

    def test_manifest_cannot_authorize_submission(self):
        path = self.folder / '07_MANIFEST.json'
        data = json.loads(path.read_text()); data['submission_authorized'] = True
        a.write_json(path, data)
        with self.assertRaisesRegex(a.AssemblyError, 'authorizing'): a.verify(self.folder)

    def test_manifest_cannot_omit_file(self):
        path = self.folder / '07_MANIFEST.json'
        data = json.loads(path.read_text()); data['files'].pop('03_FINANCIALS_MISSING.txt')
        a.write_json(path, data)
        with self.assertRaisesRegex(a.AssemblyError, 'universe'): a.verify(self.folder)

    def test_receipt_is_not_visual_review(self):
        self.assertEqual(self.receipt['visual_review'], 'NOT_ESTABLISHED_BY_BUILD')
        self.assertEqual(self.receipt['bid_status'], a.STATUS)

    def test_unsupported_glyph_not_silently_substituted(self):
        self.packet['title'] = 'A title with unsupported \u4f60\u597d'
        self.write_packet()
        out = self.root / 'unsupported'
        with self.assertRaisesRegex(a.AssemblyError, 'UNSUPPORTED_PDF_GLYPH'): a.build(self.packet_path, out)
        self.assertTrue((out / 'BUILD_INCOMPLETE.txt').exists())

    def test_html_text_is_escaped(self):
        packet, source = a.load_packet(self.packet_path)
        packet['title'] = '<script>alert(1)</script>'
        result = a.make_html(packet, a.sections(packet, source))
        self.assertNotIn('<script>', result)
        self.assertIn('&lt;script&gt;', result)

    def test_cli_verify_succeeds_with_explicit_review_status(self):
        done = subprocess.run([sys.executable, str(ROOT / 'assembler.py'), 'verify', str(self.folder)], capture_output=True, text=True)
        self.assertEqual(done.returncode, 0, done.stderr)
        self.assertFalse(json.loads(done.stdout)['submission_authorized'])

    def test_cli_wrong_output_refuses(self):
        (self.folder / 'unexpected.txt').write_text('foreign')
        done = subprocess.run([sys.executable, str(ROOT / 'assembler.py'), 'verify', str(self.folder)], capture_output=True, text=True)
        self.assertEqual(done.returncode, 2)
        self.assertIn('ASSEMBLY_REFUSED', done.stderr)

    def test_large_source_does_not_silently_overflow(self):
        source = (('Long paragraph ' * 400) + '\n\n') * 4
        path = self.inputs / 'COMMERCIAL.md'; path.write_text(source)
        self.packet['source']['sha256'] = a.digest(path.read_bytes())
        self.packet['source']['git_blob_sha'] = a.git_digest(path.read_bytes())
        self.write_packet()
        out = self.root / 'overflow'
        with self.assertRaisesRegex(a.AssemblyError, 'pagination overflow'): a.build(self.packet_path, out)
        self.assertTrue((out / 'BUILD_INCOMPLETE.txt').exists())

    def test_input_bytes_are_never_modified(self):
        before = {p.name: p.read_bytes() for p in self.inputs.iterdir()}
        a.build(self.packet_path, self.root / 'new')
        self.assertEqual(before, {p.name: p.read_bytes() for p in self.inputs.iterdir()})


if __name__ == '__main__':
    unittest.main()
