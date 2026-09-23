"""Portable reader projections. DOCX/PDF are not lossless import formats.

Every field remains traceable by JSON Pointer. Canonical JSON is the machine
record; rendered whitespace and pagination must not be used to reconstruct it.
"""
from __future__ import annotations

import argparse
from collections import OrderedDict
from pathlib import Path
import re
from typing import Any
from xml.sax.saxutils import escape

try:
    from . import transport as t
    from .workbook import display_rows
except ImportError:
    import transport as t
    from workbook import display_rows

TITLE = "Evidence interchange rehearsal"
NOTICE = "SYNTHETIC PREPARATION / NOT A UNIVERSITY FINDING"
INTRO = (
    "This reader copy demonstrates evidence and recommendation handoffs. It does not "
    "establish source authenticity, currentness, maturity, acceptance or permission to act. "
    "The accompanying JSON is the machine-readable record; CSV and XLSX can recover that "
    "record exactly. DOCX and PDF are human-readable projections, not import formats."
)
LEGEND = (
    "NULL means a field is present with no value. EMPTY STRING means present text of length "
    "zero. An absent member has no row; container counts preserve that distinction. "
    "Text beginning with =, +, - or @ is displayed as text and is not evaluated. "
    "Multiline/control characters use explicit JSON escapes in the reader projection. "
    "Dates and time-zone offsets are retained literally."
)


def sections(document: Any) -> list[tuple[str, list[list[str]]]]:
    grouped: dict[str, list[list[str]]] = OrderedDict()
    for row in display_rows(document)[1:]:
        pointer = row[0][len("Path: "):]
        match = re.match(r"^/(records|cells|sources)/([0-9]+)(?:/|$)", pointer)
        title = f"{match[1].capitalize()} / {int(match[2]) + 1}" if match else "Document envelope"
        grouped.setdefault(title, []).append(row)
    return list(grouped.items())


def write_docx(document: Any, path: str | Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    destination = Path(path)
    if destination.exists():
        raise FileExistsError(destination)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "DejaVu Sans")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06
    for name, size in [("Title", 24), ("Heading 1", 15), ("Heading 2", 10)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string("17324D")
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(5)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(2)
    header = section.header.paragraphs[0]
    header.text = "TJLabs  /  UIOWA-096  /  Reader copy"
    header.style = doc.styles["Header"]
    footer = section.footer.paragraphs[0]
    footer.text = "Synthetic rehearsal · Format integrity only  |  Page "
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_heading(TITLE, 0)
    doc.add_paragraph(NOTICE, "Subtitle")
    doc.add_paragraph(INTRO)
    doc.add_heading("How to read the fields", 1)
    doc.add_paragraph(LEGEND)
    doc.add_heading("Canonical JSON SHA-256", 2)
    doc.add_paragraph(t.document_sha256(document))
    for i, (title, rows) in enumerate(sections(document)):
        if i:
            doc.add_page_break()
        doc.add_heading(title, 1)
        for location, kind, value in rows:
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(location + "  [" + kind + "]")
            doc.add_paragraph(value)
    # Remove unused template styles (not document content) to keep the sample
    # package compact. Retain all base/link dependencies of the used styles.
    used = {p.style.style_id for p in doc.paragraphs}
    used.update({"Header", "Footer", "DefaultParagraphFont", "TableNormal", "NoList"})
    changed = True
    while changed:
        changed = False
        for style in list(doc.styles):
            if style.style_id in used:
                for tag in ("basedOn", "link", "next"):
                    for item in style.element.findall(qn("w:" + tag)):
                        target = item.get(qn("w:val"))
                        if target and target not in used:
                            used.add(target); changed = True
    for style in list(doc.styles):
        if style.style_id not in used:
            style.delete()
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/stylesWithEffects"):
            doc.part.drop_rel(rid)
    for rid, rel in list(doc.part.package.rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            del doc.part.package.rels[rid]
    with destination.open("xb") as output:
        doc.save(output)


def _font_path(font_path: str | Path | None) -> Path:
    choices = [Path(font_path)] if font_path else [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf"),
        Path("C:/Windows/Fonts/arial.ttf")]
    for choice in choices:
        if choice.is_file():
            return choice
    raise t.InterchangeError("PDF requires a Unicode TrueType font; pass --font /path/font.ttf")


def write_pdf(document: Any, path: str | Path, font_path: str | Path | None = None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    font = _font_path(font_path)
    font_name = "TessellaUnicode"
    registered_font = TTFont(font_name, str(font))
    required_text = TITLE + NOTICE + INTRO + LEGEND + t.document_sha256(document)
    required_text += "".join("".join(row) for _, rows in sections(document) for row in rows)
    missing = sorted({ord(c) for c in required_text if not c.isspace()
                      and ord(c) not in registered_font.face.charToGlyph})
    if missing:
        raise t.InterchangeError("PDF font lacks required glyphs: " +
                                 ", ".join(f"U+{cp:04X}" for cp in missing[:12]) +
                                 "; supply a font covering the source text")
    pdfmetrics.registerFont(registered_font)
    navy = colors.HexColor("#17324D")
    styles = {
        "title": ParagraphStyle("title", fontName=font_name, fontSize=23, leading=29, textColor=navy, spaceAfter=12),
        "heading": ParagraphStyle("heading", fontName=font_name, fontSize=14, leading=18, textColor=navy, spaceBefore=10, spaceAfter=7, keepWithNext=True),
        "label": ParagraphStyle("label", fontName=font_name, fontSize=9, leading=12, textColor=navy, spaceBefore=5, spaceAfter=2, keepWithNext=True),
        "body": ParagraphStyle("body", fontName=font_name, fontSize=9.5, leading=13, spaceAfter=5, splitLongWords=True),
    }
    story = []
    def add(text: str, kind: str = "body") -> None:
        story.append(Paragraph(escape(text).replace("\n", "<br/>"), styles[kind]))
    add(TITLE, "title"); add(NOTICE); add(INTRO)
    add("How to read the fields", "heading"); add(LEGEND)
    add("Canonical JSON SHA-256", "label"); add(t.document_sha256(document))
    for i, (title, rows) in enumerate(sections(document)):
        if i:
            story.append(PageBreak())
        add(title, "heading")
        for location, kind, value in rows:
            add(location + "  [" + kind + "]", "label")
            add(value)
    def page(canvas, doc):
        canvas.saveState(); canvas.setFont(font_name, 8)
        canvas.setFillColor(navy)
        canvas.drawString(54, 760, "TJLabs  /  UIOWA-096  /  Reader copy")
        canvas.drawString(54, 28, "Synthetic rehearsal · Format integrity only")
        canvas.drawRightString(558, 28, f"Page {doc.page}")
        canvas.restoreState()
    with Path(path).open("xb") as output:
        pdf = SimpleDocTemplate(output, pagesize=letter, leftMargin=54, rightMargin=54,
                                topMargin=50, bottomMargin=48, title=TITLE, author="TJLabs")
        pdf.build(story, onFirstPage=page, onLaterPages=page)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("--font", type=Path)
    args = parser.parse_args(argv)
    try:
        document = t.read_json(args.source)
        if args.destination.suffix.lower() == ".docx":
            write_docx(document, args.destination)
        elif args.destination.suffix.lower() == ".pdf":
            write_pdf(document, args.destination, args.font)
        else:
            raise t.InterchangeError("destination must end in .docx or .pdf")
        print("PASS reader projection; canonical JSON SHA-256 " + t.document_sha256(document))
    except (OSError, t.InterchangeError) as exc:
        parser.exit(2, f"reader: {exc}\n")


if __name__ == "__main__":
    main()
